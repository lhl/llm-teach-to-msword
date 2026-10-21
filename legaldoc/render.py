"""Render an IR to .docx.

Every paragraph gets exactly one thing from us: a style name. We never touch
spacing, indents, margins, fonts, or page size. That is the entire point --
it makes the class of bug the lawyer is seeing unrepresentable.

Emphasis is applied here, from the spec, by interval arithmetic. The model
declares *what* is a case name; this decides *how* a case name looks, once,
for the whole document.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as open_docx

from .ir import (
    BlockQuote,
    Caption,
    CertificateOfService,
    Document,
    Heading,
    Paragraph,
    SignatureBlock,
)
from .spec import DocSpec
from .typography import apply_sentence_spacing


class UnmatchedMark(ValueError):
    """An emphasis span the model declared does not exist in the text."""


def _style_for(block) -> str:
    return {
        "heading": lambda b: f"Heading {b.level}",
        "paragraph": lambda b: "Body Text",
        "block_quote": lambda b: "Block Quote",
        "caption": lambda b: "Caption",
        "signature": lambda b: "Signature",
        "certificate": lambda b: "Certificate",
    }[block.role](block)


def _intervals(text: str, block: Paragraph, spec: DocSpec) -> list[tuple[int, int, str]]:
    """(start, end, kind) for every declared span. Raises if one is absent."""
    wanted: list[tuple[str, str]] = []
    wanted += [(t, "b") for t in block.emphasis]
    wanted += [(t, "b") for t in block.defined_terms]
    wanted += [(t, "i") for t in block.italic]
    wanted += [(t, "i" if spec.case_name_style == "italic" else "u") for t in block.case_names]

    spans = []
    for term, kind in wanted:
        if not term:
            continue
        start = text.find(term)
        if start < 0:
            raise UnmatchedMark(
                f"declared span {term!r} does not appear in the paragraph text"
            )
        spans.append((start, start + len(term), kind))
    return spans


def _format_map(text: str, spans: list[tuple[int, int, str]]) -> list[set[str]]:
    """Per-character set of format kinds. Handles nesting and overlap."""
    fm: list[set[str]] = [set() for _ in text]
    for start, end, kind in spans:
        for i in range(start, min(end, len(text))):
            fm[i].add(kind)
    return fm


def _add_runs(par, text: str, spans: list[tuple[int, int, str]]) -> None:
    """Emit runs, coalescing adjacent characters with identical formatting."""
    if not spans:
        par.add_run(text)
        return
    fm = _format_map(text, spans)
    i = 0
    while i < len(text):
        j = i
        while j < len(text) and fm[j] == fm[i]:
            j += 1
        run = par.add_run(text[i:j])
        run.bold = "b" in fm[i] or None
        run.italic = "i" in fm[i] or None
        run.underline = "u" in fm[i] or None
        i = j


def _spaced(text: str, spec: DocSpec) -> str:
    return apply_sentence_spacing(text, spec.sentence_spaces)


def render(doc_ir: Document, spec: DocSpec, template: Path, out: Path) -> Path:
    doc = open_docx(template)

    # Start from the template's styles but an empty body.
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("}p") or child.tag.endswith("}tbl"):
            body.remove(child)

    for block in doc_ir.blocks:
        if isinstance(block, Caption):
            lines = [block.court, "", *block.parties, ""]
            lines += [block.title, "", block.document_title, block.case_number]
            for line in lines:
                doc.add_paragraph(_spaced(line, spec), style="Caption")
        elif isinstance(block, Heading):
            doc.add_paragraph(_spaced(block.text, spec), style=f"Heading {block.level}")
        elif isinstance(block, Paragraph):
            text = _spaced(block.text, spec)
            spans = _intervals(text, block, spec)
            par = doc.add_paragraph(style="Body Text")
            _add_runs(par, text, spans)
        elif isinstance(block, BlockQuote):
            doc.add_paragraph(_spaced(block.text, spec), style="Block Quote")
            if block.attribution:
                doc.add_paragraph(_spaced(block.attribution, spec), style="Block Quote")
        elif isinstance(block, SignatureBlock):
            doc.add_paragraph("Respectfully submitted,", style="Signature")
            doc.add_paragraph("", style="Signature")
            for line in [block.name, *block.address, block.phone or "", block.email or ""]:
                if line:
                    doc.add_paragraph(_spaced(line, spec), style="Signature")
        elif isinstance(block, CertificateOfService):
            doc.add_paragraph("CERTIFICATE OF SERVICE", style="Heading 1")
            doc.add_paragraph(_spaced(block.text, spec), style="Certificate")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
